"""Create reproducible invoice baskets from UCI Online Retail.

Run: python prepare_data.py [--source path/to/Online Retail.xlsx]
The small JSON output is committed so the static demo works without Python.
"""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from datetime import datetime
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
from urllib.request import Request, urlopen
from zipfile import ZipFile

from openpyxl import load_workbook


SOURCE_URL = "https://archive.ics.uci.edu/static/public/352/online%2Bretail.zip"
DATASET_URL = "https://archive.ics.uci.edu/dataset/352/online+retail"
OUTPUT = Path(__file__).resolve().parent / "data" / "baskets.json"


def read_workbook(source: Path | None) -> tuple[bytes, str | None]:
    if source is None:
        print(f"Downloading {SOURCE_URL}", flush=True)
        with urlopen(Request(SOURCE_URL, headers={"User-Agent": "Week4-Association-Rules/1.0"})) as response:
            raw = response.read()
        source_name = SOURCE_URL
    else:
        raw = source.read_bytes()
        source_name = str(source)

    archive_hash = None
    if source_name.lower().endswith(".zip"):
        archive_hash = sha256(raw).hexdigest()
        with ZipFile(BytesIO(raw)) as archive:
            xlsx_names = [name for name in archive.namelist() if name.lower().endswith(".xlsx")]
            if len(xlsx_names) != 1:
                raise ValueError(f"Expected one workbook in archive, found {xlsx_names}")
            raw = archive.read(xlsx_names[0])
    elif not source_name.lower().endswith(".xlsx"):
        raise ValueError("Source must be the official .zip or an .xlsx workbook")
    return raw, archive_hash


def create_baskets(workbook: bytes, archive_hash: str | None) -> dict:
    sheet = load_workbook(BytesIO(workbook), read_only=True, data_only=True).active
    rows = sheet.iter_rows(values_only=True)
    header = [str(value).strip() if value is not None else "" for value in next(rows)]
    columns = {name: header.index(name) for name in (
        "InvoiceNo", "StockCode", "Description", "Quantity", "InvoiceDate", "UnitPrice"
    )}

    invoices: dict[str, set[str]] = defaultdict(set)
    dates: dict[str, str] = {}
    names: dict[str, Counter] = defaultdict(Counter)
    counts = Counter()

    for row in rows:
        counts["source_rows"] += 1
        invoice = str(row[columns["InvoiceNo"]] or "").strip()
        stock = str(row[columns["StockCode"]] or "").strip()
        description = str(row[columns["Description"]] or "").strip()
        quantity = row[columns["Quantity"]]
        price = row[columns["UnitPrice"]]
        date = row[columns["InvoiceDate"]]

        # C-prefixed invoices reverse a purchase, so they are not market baskets.
        if not invoice or invoice.upper().startswith("C"):
            counts["cancelled_or_missing_invoice"] += 1
            continue
        # A purchase basket needs a product, a name and a positive sold quantity/price.
        if not stock or not description:
            counts["missing_product_or_name"] += 1
            continue
        if not isinstance(quantity, (int, float)) or quantity <= 0:
            counts["nonpositive_quantity"] += 1
            continue
        if not isinstance(price, (int, float)) or price <= 0:
            counts["nonpositive_price"] += 1
            continue
        if not isinstance(date, datetime):
            counts["missing_date"] += 1
            continue

        invoices[invoice].add(stock)
        names[stock][description] += 1
        day = date.date().isoformat()
        if invoice not in dates or day < dates[invoice]:
            dates[invoice] = day
        counts["kept_rows"] += 1

    ordered = sorted(invoices, key=lambda invoice: (dates[invoice], invoice))
    baskets = [sorted(invoices[invoice]) for invoice in ordered]
    labels = {stock: names[stock].most_common(1)[0][0] for stock in sorted(names)}
    return {
        "source": {
            "name": "UCI Online Retail",
            "url": DATASET_URL,
            "download_url": SOURCE_URL,
            "doi": "10.24432/C5BW33",
            "license": "CC BY 4.0",
            "creator": "Daqing Chen",
            "archive_sha256": archive_hash,
            "workbook_sha256": sha256(workbook).hexdigest(),
        },
        "preprocessing": {
            "method": "Each InvoiceNo is one basket; StockCode is an item; duplicate lines per invoice are collapsed. Cancellations, nonpositive quantity or price, and rows missing product/name/date are excluded.",
            **dict(sorted(counts.items())),
            "basket_count": len(baskets),
            "item_count": len(labels),
            "first_date": dates[ordered[0]],
            "last_date": dates[ordered[-1]],
        },
        "items": labels,
        "baskets": baskets,
        "dates": [dates[invoice] for invoice in ordered],
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, help="Local official .zip or .xlsx (otherwise download from UCI)")
    parser.add_argument("--output", type=Path, default=OUTPUT)
    args = parser.parse_args()
    workbook, archive_hash = read_workbook(args.source)
    data = create_baskets(workbook, archive_hash)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(json.dumps(data["preprocessing"], indent=2))
    print(f"Wrote {args.output} ({args.output.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
